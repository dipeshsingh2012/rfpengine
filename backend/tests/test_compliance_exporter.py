import io
import pytest
from starlette.testclient import TestClient

import openpyxl
import docx

from app.main import app
from app.services.compliance_exporter_service import ComplianceExporterService
from app.models.schemas import ExportRequestPayload, ExportItemPayload


@pytest.fixture
def sample_payload():
    return ExportRequestPayload(
        title="Enterprise Security Assessment",
        tenant_id="acme-corp",
        format="xlsx",
        items=[
            ExportItemPayload(
                question_index=0,
                section="Data Security",
                question_text="Are customer databases encrypted at rest using AES-256?",
                answer_text="Yes, all PostgreSQL databases and backup volumes are encrypted at rest with AWS KMS managed AES-256 keys.",
                review_status="Approved",
                assigned_role="Security Director",
                confidence_score=0.96,
                sources=[{"source_file": "01_Security_and_Compliance_Whitepaper.md", "title": "Encryption Standard"}],
                comments="Fully verified by SecOps"
            ),
            ExportItemPayload(
                question_index=1,
                section="Disaster Recovery",
                question_text="What are your RPO and RTO commitments for critical services?",
                answer_text="Our Recovery Point Objective (RPO) is <1 hour and Recovery Time Objective (RTO) is <4 hours.",
                review_status="In Review",
                assigned_role="Technical Architect",
                confidence_score=0.91,
                sources=[{"source_file": "02_SLA_Disaster_Recovery_and_Operations.pdf"}],
                comments=""
            ),
        ]
    )


def test_export_excel(sample_payload):
    sample_payload.format = "xlsx"
    content, mime, filename = ComplianceExporterService.export(sample_payload)

    assert mime == "application/vnd.openxmlformats-officedocument.spreadsheetml.sheet"
    assert filename.endswith(".xlsx")
    assert len(content) > 1000

    # Inspect openpyxl workbook
    wb = openpyxl.load_workbook(io.BytesIO(content))
    assert "Compliance Matrix" in wb.sheetnames
    assert "Audit Trail & Citations" in wb.sheetnames

    ws = wb["Compliance Matrix"]
    headers = [cell.value for cell in ws[1]]
    assert "Question / RFP Requirement" in headers
    assert "Approved Vendor Response" in headers
    assert "Review Status" in headers

    # Check question rows
    row2 = [cell.value for cell in ws[2]]
    assert "Q-001" in row2
    assert "Data Security" in row2
    assert "Are customer databases encrypted at rest using AES-256?" in row2
    assert "Approved" in row2


def test_export_docx(sample_payload):
    sample_payload.format = "docx"
    content, mime, filename = ComplianceExporterService.export(sample_payload)

    assert mime == "application/vnd.openxmlformats-officedocument.wordprocessingml.document"
    assert filename.endswith(".docx")
    assert len(content) > 1000

    doc = docx.Document(io.BytesIO(content))
    text_content = "\n".join(p.text for p in doc.paragraphs)
    assert "Enterprise Security Assessment" in text_content
    assert "Are customer databases encrypted at rest using AES-256?" in text_content
    assert "Disaster Recovery" in text_content


def test_export_pdf(sample_payload):
    sample_payload.format = "pdf"
    content, mime, filename = ComplianceExporterService.export(sample_payload)

    assert mime == "application/pdf"
    assert filename.endswith(".pdf")
    assert len(content) > 500
    assert content.startswith(b"%PDF")


def test_export_csv(sample_payload):
    sample_payload.format = "csv"
    content, mime, filename = ComplianceExporterService.export(sample_payload)

    assert mime == "text/csv; charset=utf-8"
    assert filename.endswith(".csv")
    csv_text = content.decode("utf-8-sig")
    assert "Item #" in csv_text
    assert "Data Security" in csv_text
    assert "Are customer databases encrypted at rest using AES-256?" in csv_text


def test_export_api_endpoint(sample_payload):
    client = TestClient(app)
    
    # Test XLSX export
    res_xlsx = client.post("/api/v1/responses/export", json=sample_payload.model_dump())
    assert res_xlsx.status_code == 200
    assert "spreadsheetml" in res_xlsx.headers["content-type"]
    assert len(res_xlsx.content) > 1000

    # Test PDF export
    sample_payload.format = "pdf"
    res_pdf = client.post("/api/v1/responses/export", json=sample_payload.model_dump())
    assert res_pdf.status_code == 200
    assert res_pdf.headers["content-type"] == "application/pdf"
    assert res_pdf.content.startswith(b"%PDF")

    # Test invalid format validation
    invalid_data = sample_payload.model_dump()
    invalid_data["format"] = "exe"
    res_invalid = client.post("/api/v1/responses/export", json=invalid_data)
    assert res_invalid.status_code == 422  # Pydantic schema validation error

